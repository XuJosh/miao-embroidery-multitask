"""
Translate an English PDF to Chinese (EN -> ZH) using Helsinki-NLP/opus-mt-en-zh,
then save the result as a Word document.

Usage:
    python translate_pdf_to_docx.py \
        --input papers/Zhao_ResNet50_CBAM_npj_Heritage_Science_2026.pdf \
        --output papers/Zhao_ResNet50_CBAM_中文译稿.docx
"""
import argparse
import os
import re
import sys

import pdfplumber
import torch
from docx import Document
from transformers import AutoTokenizer, pipeline


def split_text_into_sentences(text):
    """Simple English sentence splitter."""
    # Keep punctuation with the sentence.
    sentences = re.findall(r'[^.!?\n]+[.!?]*', text)
    return [s.strip() for s in sentences if s.strip()]


def chunk_paragraph(paragraph, tokenizer, max_tokens=400):
    """Split a paragraph into token-limited chunks, respecting sentence boundaries."""
    sentences = split_text_into_sentences(paragraph)
    chunks = []
    current = []
    current_len = 0
    for sent in sentences:
        # Count tokens for this sentence.
        tok_len = len(tokenizer.encode(sent, add_special_tokens=False))
        if tok_len > max_tokens:
            # Sentence alone too long: force-split by tokens.
            tokens = tokenizer.encode(sent, add_special_tokens=False)
            start = 0
            while start < len(tokens):
                end = min(start + max_tokens, len(tokens))
                chunk_tokens = tokens[start:end]
                chunks.append(tokenizer.decode(chunk_tokens, skip_special_tokens=True))
                start = end
            continue
        if current_len + tok_len > max_tokens and current:
            chunks.append(' '.join(current))
            current = [sent]
            current_len = tok_len
        else:
            current.append(sent)
            current_len += tok_len
    if current:
        chunks.append(' '.join(current))
    return chunks


def translate_texts(texts, translator, batch_size=8):
    """Translate a list of non-empty text strings."""
    results = []
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        outputs = translator(batch, max_length=512, batch_size=len(batch))
        results.extend([o['translation_text'] for o in outputs])
    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=str, required=True, help='Input PDF path')
    parser.add_argument('--output', type=str, required=True, help='Output DOCX path')
    parser.add_argument('--max_tokens', type=int, default=400)
    parser.add_argument('--batch_size', type=int, default=8)
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f'Input file not found: {args.input}')
        sys.exit(1)

    device = 0 if torch.cuda.is_available() else -1
    model_name = 'Helsinki-NLP/opus-mt-en-zh'
    print(f'Loading translation model: {model_name} (device={device})')
    translator = pipeline('translation', model=model_name, tokenizer=model_name, device=device)
    tokenizer = AutoTokenizer.from_pretrained(model_name)

    print(f'Extracting text from: {args.input}')
    doc = Document()
    doc.add_heading('ResNet50-CBAM 亚洲民族刺绣纹样分类（中文机译稿）', level=0)
    doc.add_paragraph('说明：本文档由 Helsinki-NLP/opus-mt-en-zh 模型对英文 PDF 进行机器翻译，仅供参考。')

    with pdfplumber.open(args.input) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ''
            text = text.strip()
            if not text:
                continue

            doc.add_heading(f'第 {page_idx} 页', level=1)
            # Split page text into paragraphs by blank lines.
            paragraphs = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]

            for para in paragraphs:
                chunks = chunk_paragraph(para, tokenizer, max_tokens=args.max_tokens)
                if not chunks:
                    continue
                translated_chunks = translate_texts(chunks, translator, batch_size=args.batch_size)
                translated_para = ''.join(translated_chunks)
                doc.add_paragraph(translated_para)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or '.', exist_ok=True)
    doc.save(args.output)
    print(f'Saved translated document to: {args.output}')


if __name__ == '__main__':
    main()
